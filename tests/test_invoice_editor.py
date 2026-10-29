from __future__ import annotations
import copy
import tempfile
import unittest
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
import aiosqlite
from aiogram import Bot, Dispatcher
from aiogram.client.session.base import BaseSession
from aiogram.exceptions import TelegramNotFound
from aiogram.types import Message, Chat, User, CallbackQuery, Update
from db import init_db
from db.invoices import create_invoice, get_invoice_full, save_invoice_analysis, approve_invoice, list_invoice_items, mark_invoice_analysis_processing
from services.invoice_editor import parse_value, apply_change, revision
from services.telegram_invoice_ui import display
from handlers.admin.invoice_editor import overview, cb


def analysis():
    return {'supplier': 'ООО <Компания> & партнёры', 'currency': 'RUB', 'amount_payable': 180,
            'total_amount': 180, 'calculated_total': 180, 'items': [
        {'product_name': 'Товар <1>', 'quantity': 2, 'unit_price': 90, 'line_total': 180,
         'unit_price_before_discount': 100, 'discount_percent': 10, 'discount_amount': 20, 'article': None, 'unit': 'шт'}]}


class RulesTests(unittest.TestCase):
    def test_manual_totals_include_every_item_and_repair_legacy_auto_total(self):
        from services.invoice_editor import recalculate,manual_item
        items=[manual_item({'product_name':'A','quantity':4,'unit_price_before_discount':23,'unit_price':23,'line_total':92}),
               manual_item({'product_name':'B','quantity':4,'unit_price_before_discount':400,'unit_price':400,'line_total':1600})]
        data={'recognition_source':'manual','items':items,'subtotal_before_discount':92,'amount_payable':92,'discount_amount':0}
        recalculate(data)
        self.assertEqual(data['calculated_total'],1692)
        self.assertEqual(data['subtotal_before_discount'],1692)
        self.assertEqual(data['amount_payable'],1692)
        self.assertEqual(data['total_amount'],1692)
        self.assertEqual(data['discount_amount'],0)
        changed=apply_change(data,1,'quantity',2)
        self.assertEqual(changed['amount_payable'],892)
        self.assertEqual(changed['calculated_total'],892)

    def test_manual_line_and_overall_discount_are_not_subtracted_twice(self):
        from services.invoice_editor import manual_analysis,manual_item
        items=[manual_item({'product_name':'A','quantity':2,'unit_price_before_discount':100,'unit_price':90,'line_total':180}),
               manual_item({'product_name':'B','quantity':3,'unit_price_before_discount':50,'unit_price':40,'line_total':120})]
        data=manual_analysis('Supplier',items,300)
        self.assertEqual(data['subtotal_before_discount'],350)
        self.assertEqual(data['amount_payable'],300)
        self.assertEqual(data['discount_amount'],50)
        overall=apply_change(data,-1,'discount_amount',70)
        self.assertEqual(overall['amount_payable'],280)
        self.assertEqual(overall['calculated_total'],300)
        self.assertEqual(sum(x['discount_amount'] for x in overall['items']),50)
        with self.assertRaises(ValueError): apply_change(data,-1,'discount_amount',351)
        printed=analysis()
        printed['subtotal_before_discount']=200
        printed=apply_change(printed,-1,'discount_amount',30)
        self.assertEqual(printed['amount_payable'],170)
        self.assertEqual(printed['total_amount'],170)

    def test_manual_rounding_is_per_line_and_explicit_payable_is_preserved(self):
        from services.invoice_editor import manual_analysis,manual_item
        items=[manual_item({'product_name':str(i),'quantity':0.333333,'unit_price_before_discount':0.01,'unit_price':0.01,'line_total':0}) for i in range(3)]
        tiny=manual_analysis('Supplier',items,0)
        self.assertEqual(tiny['subtotal_before_discount'],0)
        item=manual_item({'product_name':'A','quantity':2,'unit_price_before_discount':100,'unit_price':90,'line_total':180})
        fixed=manual_analysis('Supplier',[item],170)
        changed=apply_change(fixed,0,'quantity',3)
        self.assertEqual(changed['calculated_total'],270)
        self.assertEqual(changed['amount_payable'],170)
        self.assertEqual(changed['subtotal_before_discount'],300)
        self.assertEqual(changed['discount_amount'],130)

    def test_numbers_dates_and_nonfinite(self):
        self.assertEqual(parse_value('unit_price', '1 250,50'), 1250.5)
        self.assertEqual(parse_value('invoice_date', '29.02.2024'), '2024-02-29')
        for field, value in [('unit_price','nan'), ('quantity','0'), ('quantity','0.00000001'), ('quantity','-1'), ('discount_percent','101'), ('invoice_date','29.02.2025')]:
            with self.assertRaises(ValueError): parse_value(field, value)

    def test_calculations_preserve_printed_total_and_other_fields(self):
        a = analysis()
        updated = apply_change(a, 0, 'quantity', 3)
        self.assertEqual(updated['items'][0]['line_total'], 270)
        self.assertEqual(updated['calculated_total'], 270)
        self.assertEqual(updated['amount_payable'], 180)
        self.assertEqual(a['items'][0]['quantity'], 2)
        updated = apply_change(updated, 0, 'discount_percent', 20)
        self.assertEqual(updated['items'][0]['line_total'], 240)
        self.assertEqual(updated['items'][0]['unit_price'], 80)
        updated = apply_change(updated, 0, 'line_total', 150)
        self.assertEqual(updated['items'][0]['unit_price'], 50)
        self.assertEqual(updated['items'][0]['discount_amount'], 150)

    def test_mobile_tables_have_two_columns_and_full_row_actions(self):
        import json
        from html.parser import HTMLParser
        from services.telegram_invoice_ui import cell_text, embed_controls
        from handlers.admin.invoices_panel import panel_view
        class Layout(HTMLParser):
            def __init__(self):
                super().__init__(); self.width=0; self.widths=[]; self.cell={}; self.action_spans=[]
            def handle_starttag(self,tag,attrs):
                attrs=dict(attrs)
                if tag=='tr': self.width=0
                if tag in {'td','th'}:
                    self.width+=int(attrs.get('colspan',1)); self.cell=attrs
                if tag=='tg-button' and self.cell:
                    self.action_spans.append(int(self.cell.get('colspan',1)))
            def handle_endtag(self,tag):
                if tag=='tr': self.widths.append(self.width)
                if tag in {'td','th'}: self.cell={}
        a=analysis()
        inv={'id':123, 'status':'pending', 'analysis_status':'completed', 'analysis_json':json.dumps(a), 'document_total':5149200}
        _,kb,rich=overview(inv,a)
        _,panel_kb,panel=panel_view('pending',[inv],0,1,1,{'pending':1,'approved':0,'rejected':0})
        for content in (embed_controls(rich,kb),embed_controls(panel,panel_kb)):
            parser=Layout(); parser.feed(content)
            self.assertTrue(all(n<=2 for n in parser.widths))
            self.assertTrue(all(n==2 for n in parser.action_spans))
            self.assertNotIn('\u00a0',content)
        raw='A'*80+'<&'
        wrapped=cell_text(raw)
        self.assertIn('\u200b',wrapped)
        self.assertIn('&lt;&amp;',wrapped)

    def test_large_invoice_navigation_and_escaping(self):
        a = analysis(); a['items'] *= 150
        inv={'id': 123, 'status':'pending', 'analysis_json':'{}', 'analysis_status':'completed'}
        for page in (0,1,24,100):
            text,kb,rich=overview(inv,a,page=page)
            self.assertLess(len(text),4096)
            self.assertNotIn('ООО <Компания>',text)
            self.assertIn('<table bordered striped compact>',rich)
            self.assertIn('<tg-button',rich)
            self.assertEqual(sum(1 for row in kb.inline_keyboard for b in row if ':item:' in b.callback_data),6)
            for row in kb.inline_keyboard:
                for b in row: self.assertLessEqual(len(b.callback_data.encode()),64)
        _,kb,_=overview(inv,a,admin=False)
        self.assertFalse(any(':add:' in b.callback_data or ':confirm:' in b.callback_data for row in kb.inline_keyboard for b in row))


class FakeSession(BaseSession):
    def __init__(self):
        super().__init__(); self.calls=[]; self.seq=100
    async def close(self): pass
    async def stream_content(self,*args,**kwargs):
        if False: yield b''
    async def make_request(self,bot,method,timeout=None):
        self.calls.append(method); self.seq+=1
        if method.__api_method__ == 'answerCallbackQuery': return True
        return Message(message_id=self.seq,date=datetime.now(timezone.utc),chat=Chat(id=int(getattr(method,'chat_id',1)),type='private'),text=getattr(method,'text',None) or 'Card')


class EditorFlowTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        from handlers.admin.invoice_editor import router as editor
        from handlers.admin.invoice_manual import router as manual
        from handlers.admin.invoices_panel import router as panel
        from handlers.admin.invoices import router as invoices
        from handlers.user.invoices import router as user
        cls.dp=Dispatcher()
        for router in (manual,editor,panel,invoices,user): cls.dp.include_router(router)

    async def asyncSetUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.path=self.tmp.name+'/test.sqlite'
        await init_db(self.path)
        async with aiosqlite.connect(self.path) as db:
            await db.execute("INSERT INTO users (tg_id,status,created_at,updated_at) VALUES (2,'approved','now','now')")
            await db.commit()
        self.ident=await create_invoice(self.path,2,'test-file','photo')
        await save_invoice_analysis(self.path,self.ident,analysis())
        self.settings=SimpleNamespace(db_path=self.path,admin_ids={1})
        self.session=FakeSession(); self.bot=Bot('123456:FAKE_TEST_TOKEN',session=self.session)
        for uid in (1,2,3): await self.dp.fsm.get_context(bot=self.bot,user_id=uid,chat_id=uid).clear()

    async def asyncTearDown(self):
        await self.bot.session.close(); self.tmp.cleanup()

    async def message(self,text,uid=1):
        msg=Message(message_id=1,date=datetime.now(timezone.utc),chat=Chat(id=uid,type='private'),from_user=User(id=uid,is_bot=False,first_name='Test'),text=text)
        await self.dp.feed_update(self.bot,Update(update_id=1,message=msg),settings=self.settings)

    async def click(self,data,uid=1):
        msg=Message(message_id=10,date=datetime.now(timezone.utc),chat=Chat(id=uid,type='private'),text='Card')
        query=CallbackQuery(id='test',from_user=User(id=uid,is_bot=False,first_name='Test'),chat_instance='test',message=msg,data=data)
        await self.dp.feed_update(self.bot,Update(update_id=2,callback_query=query),settings=self.settings)

    async def inv(self): return await get_invoice_full(self.path,self.ident)

    async def test_admin_menu_does_not_open_user_menu(self):
        await self.message('🧾 Накладные')
        self.assertIn('На проверке',self.session.calls[-1].rich_message['html'])
        self.assertIn('iv:page:',self.session.calls[-1].rich_message['html'])
        self.assertEqual(self.session.calls[-1].reply_markup.inline_keyboard, [])

    async def test_user_can_open_pending_invoices_from_menu(self):
        await create_invoice(self.path,3,'other-user-file','photo')
        await self.message('🧾 Накладные',uid=2)
        menu=self.session.calls[-1]
        pending=next(b for row in menu.reply_markup.inline_keyboard for b in row
                     if b.text=='🟡 На проверке')
        self.assertIn('На проверке: <b>1</b>',menu.text)
        await self.click(pending.callback_data,uid=2)
        listing=self.session.calls[-1]
        self.assertIn('Накладные на проверке',listing.text)
        opens=[b.callback_data for row in listing.reply_markup.inline_keyboard for b in row
               if b.callback_data.startswith('uinv:open:')]
        self.assertEqual(opens,[f'uinv:open:{self.ident}'])
        await self.click(opens[0],uid=2)
        original=self.session.calls[-1]
        self.assertEqual(original.__api_method__,'sendPhoto')
        self.assertEqual(original.photo,'test-file')
        self.assertIn('На проверке',original.caption)
        self.assertFalse(any(b.callback_data.startswith('iv:')
                             for row in original.reply_markup.inline_keyboard for b in row))
        await self.click(cb('page',await self.inv()),uid=2)
        self.assertIn('только администратору',self.session.calls[-1].text)

    async def test_pending_invoice_document_and_foreign_file_access(self):
        ident=await create_invoice(self.path,2,'document-file','document')
        await self.click(f'uinv:open:{ident}',uid=2)
        self.assertEqual(self.session.calls[-1].__api_method__,'sendDocument')
        self.assertEqual(self.session.calls[-1].document,'document-file')
        self.session.calls.clear()
        await self.click(f'uinv:open:{ident}',uid=3)
        self.assertIn('не найдена',self.session.calls[-1].text)
        self.assertFalse(any(c.__api_method__ in {'sendPhoto','sendDocument'}
                             for c in self.session.calls))

    async def test_pending_list_pagination_and_return_to_menu(self):
        expected={self.ident}
        for index in range(11):
            expected.add(await create_invoice(self.path,2,f'file-{index}','photo'))
        seen=set()
        for page in range(3):
            await self.click(f'uinv:list:pending:{page}',uid=2)
            listing=self.session.calls[-1]
            rows=listing.reply_markup.inline_keyboard
            opens=[b.callback_data for row in rows for b in row
                   if b.callback_data.startswith('uinv:open:')]
            ids={int(data.split(':')[-1]) for data in opens}
            self.assertFalse(seen & ids)
            seen.update(ids)
            self.assertEqual(len(opens),5 if page<2 else 2)
            self.assertTrue(all(len(row)==1 for row in rows[:len(opens)]))
            self.assertEqual(rows[-1][0].callback_data,'uinv:menu')
            if page==1:
                self.assertEqual([b.callback_data for b in rows[-2]],
                                 ['uinv:list:pending:0','uinv:list:pending:2'])
        self.assertEqual(seen,expected)
        await self.click('uinv:list:pending:999',uid=2)
        self.assertIn('Стр. <b>3</b>/<b>3</b>',self.session.calls[-1].text)
        await self.click('uinv:list:pending:-1',uid=2)
        self.assertIn('Стр. <b>1</b>/<b>3</b>',self.session.calls[-1].text)
        await self.click('uinv:menu',uid=2)
        self.assertIn('На проверке: <b>12</b>',self.session.calls[-1].text)

    async def test_pending_list_updates_after_approval_and_empty_list_has_menu(self):
        await approve_invoice(self.path,self.ident,180,9,analysis()['items'])
        await self.click('uinv:list:pending:0',uid=2)
        listing=self.session.calls[-1]
        self.assertIn('Накладные на проверке',listing.text)
        self.assertIn('Пока пусто',listing.text)
        self.assertEqual(listing.reply_markup.inline_keyboard[0][0].callback_data,'uinv:menu')
        await self.click('uinv:list:approved:0',uid=2)
        self.assertIn('Принятые накладные',self.session.calls[-1].text)
        self.assertIn('Вознаграждение',self.session.calls[-1].text)
        await self.click('uinv:list:rejected:0',uid=2)
        self.assertIn('Отклонённые накладные',self.session.calls[-1].text)

    async def test_invalid_user_invoice_list_callbacks_show_alert(self):
        for data in ('uinv:list:unknown:0','uinv:list:pending:no',
                     'uinv:list:pending','uinv:list:pending:0:extra'):
            with self.subTest(data=data):
                self.session.calls.clear()
                await self.click(data,uid=2)
                self.assertEqual(len(self.session.calls),1)
                self.assertTrue(self.session.calls[-1].show_alert)

    async def test_table_actions_are_unique_and_fields_are_inside_cells(self):
        from html.parser import HTMLParser
        class Inspect(HTMLParser):
            def __init__(self):
                super().__init__(); self.path=[]; self.actions=[]; self.tables=[]
            def handle_starttag(self, tag, attrs):
                attrs=dict(attrs)
                if tag=='table': self.tables.append(attrs)
                if tag=='tg-button': self.actions.append((attrs.get('data'), 'td' in self.path))
                if tag not in {'br','hr'}: self.path.append(tag)
            def handle_endtag(self, tag):
                if self.path and self.path[-1] == tag: self.path.pop()
        inv=await self.inv()
        await self.click(cb('page',inv))
        method=self.session.calls[-1]
        self.assertEqual(method.reply_markup.inline_keyboard, [])
        doc=Inspect(); doc.feed(method.rich_message['html'])
        actions=[a for a,_ in doc.actions]
        self.assertEqual(len(actions),len(set(actions)))
        self.assertTrue(all('bordered' in t and 'striped' in t for t in doc.tables))
        self.assertTrue(any(':item:' in a and inside for a,inside in doc.actions))
        await self.click(cb('item',inv,index=0))
        method=self.session.calls[-1]
        self.assertEqual(method.reply_markup.inline_keyboard, [])
        doc=Inspect(); doc.feed(method.rich_message['html'])
        fields=[(a,inside) for a,inside in doc.actions if ':field:' in a]
        self.assertEqual(len(fields),9)
        self.assertTrue(all(inside for _,inside in fields))

    async def test_native_status_switch_and_back_from_confirmation(self):
        inv=await self.inv()
        await self.click(cb('confirm',inv))
        self.assertIn('Процент вознаграждения',self.session.calls[-1].rich_message['html'])
        await self.click('ainv:cancel')
        self.assertIn('<table bordered',self.session.calls[-1].rich_message['html'])
        await self.click('ainv:panel:rejected:0')
        content=self.session.calls[-1].rich_message['html']
        self.assertIn('пока нет накладных',content)
        self.assertIn('ainv:panel:pending:0',content)
        self.assertIn('ainv:panel:approved:0',content)
        self.assertEqual(self.session.calls[-1].reply_markup.inline_keyboard, [])

    async def test_old_server_receives_full_classic_keyboard_on_edit(self):
        from services.telegram_invoice_ui import EditRichInvoice
        inv=await self.inv()
        text,kb,rich=overview(inv,analysis())
        fake=SimpleNamespace(chat=SimpleNamespace(id=1),message_id=1,bot=AsyncMock(),edit_text=AsyncMock())
        fake.bot.side_effect=TelegramNotFound(method=EditRichInvoice(chat_id=1,message_id=1,rich_message={'html':'x'}),message='Not Found')
        await display(fake,text,kb,rich=rich,edit=True)
        fake.edit_text.assert_awaited_once_with(text,parse_mode='HTML',reply_markup=kb)
        self.assertTrue(any(':item:' in b.callback_data for row in kb.inline_keyboard for b in row))

    async def test_field_edit_persists_and_cancel_does_not_modify(self):
        inv=await self.inv()
        await self.click(cb('field',inv,index=0,field=0))
        await self.message('Исправленный товар')
        updated=await self.inv(); self.assertIn('Исправленный товар',updated['analysis_json'])
        await self.click(cb('field',updated,index=0,field=0))
        await self.message('/cancel')
        self.assertEqual(updated['analysis_json'],(await self.inv())['analysis_json'])

    async def test_add_three_steps_delete_confirmation_and_stale_button(self):
        inv=await self.inv(); await self.click(cb('add',inv))
        await self.message('Новый товар'); await self.message('3'); await self.message('25,50')
        from db.invoices import decode_invoice_analysis
        inv=await self.inv(); a=decode_invoice_analysis(inv)
        self.assertEqual(len(a['items']),2); self.assertEqual(a['items'][1]['line_total'],76.5)
        await self.click(cb('delete',inv,index=1))
        self.assertEqual(len(decode_invoice_analysis(await self.inv())['items']),2)
        await self.click(cb('deleteok',inv,index=1))
        self.assertEqual(len(decode_invoice_analysis(await self.inv())['items']),1)
        await self.click(cb('deleteok',inv,index=0)) # obsolete revision must not delete another item
        self.assertEqual(len(decode_invoice_analysis(await self.inv())['items']),1)

    async def test_invoice_details_are_admin_only_even_for_owner(self):
        inv=await self.inv()
        await self.click(cb('page',inv),uid=2)
        await self.click(cb('field',inv,index=0,field=0),uid=2)
        self.assertIn('только администратору',self.session.calls[-1].text)
        await self.click(cb('page',inv),uid=3)
        self.assertIn('только администратору',self.session.calls[-1].text)
        self.assertEqual(inv['analysis_json'],(await self.inv())['analysis_json'])

    async def test_upload_only_sends_pending_status_to_manager(self):
        from aiogram.types import Document
        from handlers.user.invoices import NewInvoice
        from services.invoice_recognition import InvoiceRecognitionError
        for failure in (False, True):
            with self.subTest(recognition_failed=failure):
                context= self.dp.fsm.get_context(bot=self.bot,user_id=2,chat_id=2)
                await context.set_state(NewInvoice.waiting_file)
                self.session.calls.clear()
                message=Message(message_id=10,date=datetime.now(timezone.utc),chat=Chat(id=2,type='private'),
                                from_user=User(id=2,is_bot=False,first_name='Manager'),
                                document=Document(file_id='invoice-document',file_unique_id='invoice-document',file_name='invoice.pdf',mime_type='application/pdf'))
                with patch('handlers.user.invoices.analyze_invoice_from_telegram',new_callable=AsyncMock) as recognize, \
                     patch('handlers.user.invoices._notify_admins_new_invoice',new_callable=AsyncMock), \
                     patch('handlers.user.invoices._notify_admins_duplicate_invoice',new_callable=AsyncMock), \
                     patch('handlers.user.invoices._notify_admins_recognized',new_callable=AsyncMock):
                    if failure:
                        recognize.side_effect=InvoiceRecognitionError('Internal recognition error')
                    else:
                        recognize.return_value=analysis()
                    await self.dp.feed_update(self.bot,Update(update_id=3,message=message),settings=self.settings)
                self.assertEqual(len(self.session.calls),1)
                self.assertEqual(self.session.calls[0].__api_method__,'sendMessage')
                self.assertIn('на проверке',self.session.calls[0].text)
                self.assertNotIn('Internal recognition',self.session.calls[0].text)
                self.assertNotIn('Товар',self.session.calls[0].text)

    async def test_approval_notifies_manager_with_reward_without_decoded_data(self):
        await self.click(cb('confirm',await self.inv()))
        self.session.calls.clear()
        await self.message('5')
        manager_calls=[method for method in self.session.calls if getattr(method,'chat_id',None)==2]
        self.assertEqual(len(manager_calls),1)
        self.assertEqual(manager_calls[0].__api_method__,'sendMessage')
        self.assertIn('накладная принята',manager_calls[0].text)
        self.assertIn('9,00',manager_calls[0].text)
        self.assertNotIn('Товарных позиций',manager_calls[0].text)
        self.assertNotIn('Сумма продаж',manager_calls[0].text)
        self.assertEqual((await self.inv())['status'],'approved')

    async def test_manager_excel_report_has_no_invoice_product_breakdown(self):
        from pathlib import Path
        from openpyxl import load_workbook
        from services.analytics_report import build_user_report_xlsx,build_admin_report_xlsx
        await approve_invoice(self.path,self.ident,180,9,analysis()['items'])
        for admin in (False,True):
            if admin:
                path=await build_admin_report_xlsx(self.path,'2020-01-01','2030-01-01','custom')
            else:
                path=await build_user_report_xlsx(self.path,2,'2020-01-01','2030-01-01','custom')
            try:
                book=load_workbook(path)
                values=[cell.value for sheet in book for row in sheet for cell in row]
                self.assertEqual(analysis()['items'][0]['product_name'] in values,admin)
                if not admin:
                    self.assertNotIn('Товары',book.sheetnames)
                    self.assertNotIn('Сводка товаров',book.sheetnames)
                book.close()
            finally:
                Path(path).unlink()

    async def test_concurrent_edits_and_approval_snapshot(self):
        inv=await self.inv(); version=revision(inv)
        await save_invoice_analysis(self.path,self.ident,apply_change(analysis(),0,'quantity',3),expected_revision=version)
        with self.assertRaises(ValueError):
            await save_invoice_analysis(self.path,self.ident,apply_change(analysis(),0,'quantity',4),expected_revision=version)
        with self.assertRaises(ValueError):
            await approve_invoice(self.path,self.ident,180,9,analysis()['items'],expected_revision=version)
        self.assertEqual((await self.inv())['status'],'pending')

    def manual_state(self, uid=1):
        return self.dp.fsm.get_context(bot=self.bot,user_id=uid,chat_id=uid)

    async def manual_button(self, name):
        from handlers.admin.invoice_manual import action
        data=await self.manual_state().get_data()
        callback=action(data,name)
        await self.click(callback)
        return callback

    async def start_empty_manual(self):
        from db.invoices import fail_invoice_analysis
        self.ident=await create_invoice(self.path,2,'manual-file','photo')
        await fail_invoice_analysis(self.path,self.ident,'Recognition failed')
        await self.click(f'ainv:manual:{self.ident}')

    async def test_manual_failed_invoice_two_items_edit_and_approve(self):
        from db.invoices import decode_invoice_analysis
        await self.start_empty_manual()
        for text in ('ООО <Вручную>', 'Товар А', '2', '1000', '900', '-'):
            await self.message(text)
        self.assertIsNone((await self.inv())['analysis_json'])
        await self.manual_button('next')
        for text in ('Товар Б', '0,5', '500', '400', '190'):
            await self.message(text)
        await self.manual_button('total')
        await self.message('-')
        self.assertIsNone((await self.inv())['analysis_json'])
        await self.manual_button('save')
        inv=await self.inv(); data=decode_invoice_analysis(inv)
        self.assertEqual(data['supplier'],'ООО <Вручную>')
        self.assertEqual(data['recognition_source'],'manual')
        self.assertEqual(len(data['items']),2)
        self.assertEqual(data['items'][0]['unit_price_before_discount'],1000)
        self.assertEqual(data['items'][0]['discount_amount'],200)
        self.assertEqual(data['items'][1]['unit_price'],380)
        self.assertEqual(data['calculated_total'],1990)
        self.assertEqual(data['amount_payable'],1990)
        self.assertEqual(inv['analysis_status'],'completed')
        self.assertIsNone(inv['analysis_error'])
        self.assertEqual(inv['status'],'pending')
        self.assertIsNone(inv['reward_amount'])
        self.assertEqual(await list_invoice_items(self.path,self.ident),[])
        await self.click(cb('field',inv,index=0,field=0))
        await self.message('Исправленный вручную товар')
        inv=await self.inv()
        self.assertEqual(decode_invoice_analysis(inv)['items'][0]['product_name'],'Исправленный вручную товар')
        await self.click(cb('confirm',inv))
        await self.message('5')
        inv=await self.inv()
        self.assertEqual(inv['status'],'approved')
        self.assertEqual(inv['deal_amount'],1990)
        self.assertEqual(inv['reward_amount'],99.5)
        self.assertEqual(len(await list_invoice_items(self.path,self.ident)),2)

    async def test_manual_cancel_preserves_original_and_existing_items_require_confirmation(self):
        inv=await self.inv()
        await self.click(cb('manual',inv))
        self.assertIn('заменит',self.session.calls[-1].rich_message['html'])
        self.assertEqual((await self.manual_state().get_data())['step'],'replace')
        await self.manual_button('begin')
        await self.message('Новый поставщик')
        await self.message('/cancel')
        self.assertEqual((await self.inv())['analysis_json'],inv['analysis_json'])
        self.assertIsNone(await self.manual_state().get_state())

    async def test_manual_validation_back_and_no_discount(self):
        from db.invoices import decode_invoice_analysis
        await self.start_empty_manual()
        await self.message('-')
        self.assertEqual((await self.manual_state().get_data())['step'],'supplier')
        await self.message('Поставщик')
        await self.message('Товар')
        for text in ('0','-1','nan','inf','abc'):
            await self.message(text)
            self.assertEqual((await self.manual_state().get_data())['step'],'quantity')
        await self.message('2')
        await self.message('100')
        await self.message('101')
        self.assertEqual((await self.manual_state().get_data())['step'],'unit_price')
        await self.message('-')
        await self.message('201')
        self.assertEqual((await self.manual_state().get_data())['step'],'line_total')
        await self.manual_button('back')
        self.assertEqual((await self.manual_state().get_data())['step'],'unit_price')
        await self.message('-')
        await self.message('-')
        await self.manual_button('total')
        await self.message('190')
        self.assertIn('отличается',self.session.calls[-1].rich_message['html'])
        await self.manual_button('save')
        data=decode_invoice_analysis(await self.inv())
        self.assertEqual(data['calculated_total'],200)
        self.assertEqual(data['amount_payable'],190)
        self.assertEqual(data['items'][0]['unit_price'],100)
        self.assertEqual(data['items'][0]['discount_amount'],0)

    async def test_manual_old_buttons_and_repeated_save_do_not_duplicate(self):
        from db.invoices import decode_invoice_analysis
        await self.start_empty_manual()
        for text in ('Поставщик','Товар','1','100','90','-'):
            await self.message(text)
        old=await self.manual_button('next')
        await self.click(old)
        self.assertTrue(self.session.calls[-1].show_alert)
        await self.manual_button('back')
        await self.manual_button('total')
        await self.message('-')
        save=await self.manual_button('save')
        snapshot=(await self.inv())['analysis_json']
        await self.click(save)
        self.assertTrue(self.session.calls[-1].show_alert)
        self.assertEqual((await self.inv())['analysis_json'],snapshot)
        self.assertEqual(len(decode_invoice_analysis(await self.inv())['items']),1)

    async def test_manual_detects_external_edit_and_recognition_lock(self):
        await self.start_empty_manual()
        await self.message('Поставщик')
        updated=analysis(); updated['supplier']='Изменил другой администратор'
        await save_invoice_analysis(self.path,self.ident,updated)
        snapshot=(await self.inv())['analysis_json']
        await self.message('Устаревший товар')
        self.assertIsNone(await self.manual_state().get_state())
        self.assertEqual((await self.inv())['analysis_json'],snapshot)
        await mark_invoice_analysis_processing(self.path,self.ident)
        await self.click(f'ainv:manual:{self.ident}')
        self.assertTrue(self.session.calls[-1].show_alert)

    async def test_manual_is_admin_only_and_does_not_replace_approved_invoice(self):
        inv=await self.inv()
        for uid in (2,3):
            await self.click(f'ainv:manual:{self.ident}',uid=uid)
            self.assertTrue(self.session.calls[-1].show_alert)
            await self.click(cb('manual',inv),uid=uid)
            self.assertIn('только администратору',self.session.calls[-1].text)
        self.assertEqual((await self.inv())['analysis_json'],inv['analysis_json'])
        await approve_invoice(self.path,self.ident,180,9,analysis()['items'])
        inv=await self.inv()
        await self.click(f'ainv:manual:{self.ident}')
        self.assertTrue(self.session.calls[-1].show_alert)
        await self.click(cb('manual',inv))
        self.assertIsNone(await self.manual_state().get_state())
        self.assertEqual((await self.inv())['reward_amount'],9)

    async def test_manual_reedit_last_item_and_finish_without_ai(self):
        from db.invoices import decode_invoice_analysis
        await self.start_empty_manual()
        for text in ('Поставщик','Товар','1','100','90','-'):
            await self.message(text)
        await self.manual_button('edit_item')
        for text in ('Исправленный товар','3','100','80','-'):
            await self.message(text)
        await self.manual_button('total')
        await self.message('-')
        await self.manual_button('save')
        data=decode_invoice_analysis(await self.inv())
        self.assertEqual(len(data['items']),1)
        self.assertEqual(data['items'][0]['product_name'],'Исправленный товар')
        self.assertEqual(data['amount_payable'],240)

    async def test_manual_entry_button_and_regular_editing_available(self):
        from handlers.admin.invoices import _recognize_and_show
        from services.invoice_recognition import InvoiceRecognitionError
        inv=await self.inv()
        _,kb,_=overview(inv,analysis())
        callbacks=[b.callback_data for row in kb.inline_keyboard for b in row]
        self.assertTrue(any(':manual:' in data for data in callbacks))
        self.assertTrue(any(b.text=='✍️ Ручной ввод' for row in kb.inline_keyboard for b in row))
        self.assertTrue(any(':add:' in data for data in callbacks))
        self.assertTrue(any(':header:' in data for data in callbacks))
        progress=SimpleNamespace(edit_text=AsyncMock())
        message=SimpleNamespace(answer=AsyncMock(return_value=progress),bot=self.bot)
        with patch('handlers.admin.invoices.analyze_invoice_from_telegram',new=AsyncMock(side_effect=InvoiceRecognitionError('Error'))):
            await _recognize_and_show(message,self.settings,self.ident)
        callbacks=[b.callback_data for row in progress.edit_text.call_args.kwargs['reply_markup'].inline_keyboard for b in row]
        self.assertIn(f'ainv:manual:{self.ident}',callbacks)
        self.assertIn(f'ainv:edit:{self.ident}',callbacks)

    async def test_all_header_fields_visible_with_table_buttons_and_save(self):
        import re
        import html
        from db.invoices import decode_invoice_analysis
        from services.invoice_editor import HEADER_FIELDS
        values={'document_type':'Счёт', 'invoice_number':'Р-123', 'invoice_date':'28.10.2026',
                'supplier':'Поставщик <Новый>', 'buyer':'Покупатель', 'responsible_manager':'Менеджер',
                'currency':'rub', 'subtotal_before_discount':'250', 'discount_amount':'25',
                'amount_payable':'225', 'vat_amount':'37,5'}
        for field,text in values.items():
            with self.subTest(field=field):
                await self.click(cb('header',await self.inv(),page=1))
                method=self.session.calls[-1]
                self.assertEqual(method.reply_markup.inline_keyboard,[])
                content=method.rich_message['html']
                for label in HEADER_FIELDS.values(): self.assertIn(label,content.replace('\u200b',''))
                buttons=[html.unescape(data) for data in re.findall(r'data="([^"]+)"',content) if ':field:' in data]
                self.assertEqual(len(buttons),len(HEADER_FIELDS))
                self.assertEqual(len(set(buttons)),len(buttons))
                edit=next(data for data in buttons if int(data.split(':')[5])==list(HEADER_FIELDS).index(field))
                await self.click(edit)
                self.assertEqual(self.session.calls[-1].reply_markup.inline_keyboard,[])
                self.assertIn('Сейчас:',self.session.calls[-1].rich_message['html'])
                await self.message(text)
                data=decode_invoice_analysis(await self.inv())
                self.assertEqual(data[field],parse_value(field,text))
        self.assertEqual(data['amount_payable'],data['total_amount'])

    async def test_overview_shows_empty_header_fields_without_collapsing(self):
        from services.invoice_editor import HEADER_FIELDS
        text,kb,rich=overview({'id':999,'status':'pending','analysis_status':'failed','analysis_json':None},{'items':[]})
        self.assertNotIn('<details>',rich)
        for label in HEADER_FIELDS.values():
            self.assertIn(label,text)
            self.assertIn(label,rich)
        self.assertIn('—',text)

    async def test_add_item_and_edit_header_after_failed_recognition(self):
        import re
        import html
        from db.invoices import decode_invoice_analysis,fail_invoice_analysis
        self.ident=await create_invoice(self.path,2,'failed-file','photo')
        await fail_invoice_analysis(self.path,self.ident,'Recognition failed')
        await self.click(cb('header',await self.inv()))
        header=self.session.calls[-1]
        edit=next(html.unescape(data) for data in re.findall(r'data="([^"]+)"',header.rich_message['html'])
                  if ':field:' in data and int(data.split(':')[5])==3)
        await self.click(edit)
        await self.message('Введённый поставщик')
        await self.click(cb('add',await self.inv()))
        for text in ('Пропущенный товар','2','100'): await self.message(text)
        data=decode_invoice_analysis(await self.inv())
        self.assertEqual(data['supplier'],'Введённый поставщик')
        self.assertEqual(data['items'][0]['product_name'],'Пропущенный товар')
        self.assertEqual(data['items'][0]['line_total'],200)

    async def test_manual_replacement_preserves_document_headers_and_uses_new_supplier(self):
        from db.invoices import decode_invoice_analysis
        from services.invoice_report import invoice_report_details
        original=analysis(); original.update(invoice_number='MAN-123',invoice_date='2026-10-28',buyer='Покупатель')
        await save_invoice_analysis(self.path,self.ident,original)
        async with aiosqlite.connect(self.path) as db:
            cur=await db.execute("INSERT INTO suppliers(name,created_at,updated_at) VALUES ('Старый поставщик','now','now')")
            await db.execute('UPDATE invoices SET supplier_id=? WHERE id=?',(cur.lastrowid,self.ident))
            await db.commit()
        await self.click(cb('manual',await self.inv()))
        await self.manual_button('begin')
        for text in ('Новый поставщик','Товар','1','100','-','-'):
            await self.message(text)
        await self.manual_button('total')
        await self.message('-')
        await self.manual_button('save')
        inv=await self.inv(); data=decode_invoice_analysis(inv)
        self.assertEqual(data['invoice_number'],'MAN-123')
        self.assertEqual(data['invoice_date'],'2026-10-28')
        self.assertEqual(data['buyer'],'Покупатель')
        self.assertIsNone(inv['supplier_id'])
        self.assertEqual(invoice_report_details(inv),('Новый поставщик',100))

    async def test_manual_current_draft_cannot_save_after_another_admin_approves(self):
        await self.start_empty_manual()
        for text in ('Поставщик','Товар','1','100','-','-'):
            await self.message(text)
        await self.manual_button('total')
        await self.message('-')
        await save_invoice_analysis(self.path,self.ident,analysis())
        await approve_invoice(self.path,self.ident,180,9,analysis()['items'])
        snapshot=await self.inv()
        await self.manual_button('save')
        self.assertTrue(self.session.calls[-1].show_alert)
        self.assertEqual((await self.inv())['analysis_json'],snapshot['analysis_json'])
        self.assertEqual((await self.inv())['reward_amount'],9)

    async def test_full_manual_header_and_item_fields_edit_in_draft_without_ai(self):
        from db.invoices import decode_invoice_analysis
        from services.invoice_editor import HEADER_FIELDS,ITEM_FIELDS
        await self.start_empty_manual()
        await self.manual_button('headers')
        header_values={
            'document_type':'УПД', 'invoice_number':'MAN-ALL-123', 'invoice_date':'28.10.2026',
            'supplier':'Поставщик вручную', 'buyer':'Покупатель вручную', 'responsible_manager':'Иванов',
            'currency':'usd', 'subtotal_before_discount':'500', 'discount_amount':'50',
            'amount_payable':'450', 'vat_amount':'75',
        }
        for field,text in header_values.items():
            await self.manual_button(f'field_h_-1_{list(HEADER_FIELDS).index(field)}')
            await self.message(text)
        self.assertIsNone((await self.inv())['analysis_json'])
        await self.manual_button('resume')
        self.assertEqual((await self.manual_state().get_data())['step'],'product_name')
        for text in ('Товар','5','100','90','-'): await self.message(text)
        await self.manual_button('items_0')
        await self.manual_button('item_0')
        item_values={'product_name':'Полное название', 'article':'SKU-ALL-1', 'unit':'упак',
                     'quantity':'10', 'unit_price':'45', 'line_total':'450',
                     'unit_price_before_discount':'50', 'discount_percent':'10', 'discount_amount':'50'}
        for field,text in item_values.items():
            await self.manual_button(f'field_i_0_{list(ITEM_FIELDS).index(field)}')
            await self.message(text)
        self.assertIsNone((await self.inv())['analysis_json'])
        await self.manual_button('resume')
        await self.manual_button('total')
        await self.message('-')
        await self.manual_button('save')
        inv=await self.inv(); data=decode_invoice_analysis(inv)
        for field,text in header_values.items():
            self.assertEqual(data[field],parse_value(field,text))
        item=data['items'][0]
        self.assertEqual(set(item),set(ITEM_FIELDS))
        for field,text in item_values.items():
            self.assertEqual(item[field],parse_value(field,text))
        self.assertEqual(data['recognition_source'],'manual')
        self.assertEqual(inv['status'],'pending')
        self.assertIsNone(inv['reward_amount'])
        self.assertEqual(await list_invoice_items(self.path,self.ident),[])

    async def test_manual_review_can_edit_any_item_and_clear_optional_fields(self):
        from db.invoices import decode_invoice_analysis
        from services.invoice_editor import HEADER_FIELDS,ITEM_FIELDS
        await self.start_empty_manual()
        for text in ('Supplier','First','1','100','-','-'): await self.message(text)
        await self.manual_button('next')
        for text in ('Second','1','200','-','-'): await self.message(text)
        await self.manual_button('total')
        await self.message('-')
        await self.manual_button('items_0')
        await self.manual_button('item_0')
        await self.manual_button(f'field_i_0_{list(ITEM_FIELDS).index("article")}')
        await self.message('SKU-1')
        await self.manual_button(f'field_i_0_{list(ITEM_FIELDS).index("article")}')
        await self.message('-')
        await self.manual_button(f'field_i_0_{list(ITEM_FIELDS).index("product_name")}')
        await self.message('Corrected first')
        await self.manual_button('resume')
        self.assertEqual((await self.manual_state().get_data())['step'],'review')
        await self.manual_button('headers')
        await self.manual_button(f'field_h_-1_{list(HEADER_FIELDS).index("vat_amount")}')
        await self.message('0')
        await self.manual_button(f'field_h_-1_{list(HEADER_FIELDS).index("vat_amount")}')
        await self.message('-')
        await self.manual_button('resume')
        await self.manual_button('save')
        data=decode_invoice_analysis(await self.inv())
        self.assertEqual(data['items'][0]['product_name'],'Corrected first')
        self.assertIsNone(data['items'][0]['article'])
        self.assertEqual(data['items'][1]['product_name'],'Second')
        self.assertIsNone(data['vat_amount'])

    async def test_manual_all_schema_business_fields_have_draft_edit_controls(self):
        import re
        from services.invoice_editor import HEADER_FIELDS,ITEM_FIELDS
        from services.invoice_recognition import INVOICE_SCHEMA
        await self.start_empty_manual()
        await self.manual_button('headers')
        rendered=self.session.calls[-1].rich_message['html']
        controls=[x for x in re.findall(r'data="([^"]+)"',rendered) if ':field_h_' in x]
        self.assertEqual(len(controls),len(HEADER_FIELDS))
        self.assertEqual(len(set(controls)),len(controls))
        self.assertEqual(set(INVOICE_SCHEMA['properties'])-set(HEADER_FIELDS),{'confidence','warnings','total_amount','items'})
        self.assertEqual(set(INVOICE_SCHEMA['properties']['items']['items']['properties']),set(ITEM_FIELDS))
        self.assertEqual(self.session.calls[-1].reply_markup.inline_keyboard,[])

    async def test_manual_field_validation_and_stale_draft_buttons(self):
        from services.invoice_editor import HEADER_FIELDS
        await self.start_empty_manual()
        await self.manual_button('headers')
        field_callback=await self.manual_button(f'field_h_-1_{list(HEADER_FIELDS).index("invoice_date")}')
        await self.message('31.02.2026')
        self.assertEqual((await self.manual_state().get_data())['step'],'field')
        await self.message('28.10.2026')
        await self.click(field_callback)
        self.assertTrue(self.session.calls[-1].show_alert)
        await self.manual_button('cancel')
        self.assertIsNone((await self.inv())['analysis_json'])

    async def test_manual_add_after_review_recalculates_all_items(self):
        from db.invoices import decode_invoice_analysis
        await self.start_empty_manual()
        for text in ('Supplier','First','1','100','-','-'): await self.message(text)
        await self.manual_button('total')
        await self.message('-')
        await self.manual_button('back')
        await self.manual_button('next')
        for text in ('Second','2','200','-','-'): await self.message(text)
        await self.manual_button('total')
        await self.message('-')
        await self.manual_button('save')
        data=decode_invoice_analysis(await self.inv())
        self.assertEqual(len(data['items']),2)
        self.assertEqual(data['amount_payable'],500)
        self.assertEqual(data['subtotal_before_discount'],500)
        self.assertEqual(data['discount_amount'],0)

    async def test_add_to_saved_manual_invoice_updates_total_and_reward(self):
        from db.invoices import decode_invoice_analysis
        from services.invoice_editor import manual_analysis,manual_item
        item=manual_item({'product_name':'First','quantity':4,'unit_price_before_discount':23,'unit_price':23,'line_total':92})
        await save_invoice_analysis(self.path,self.ident,manual_analysis('Supplier',[item],92))
        await self.click(cb('add',await self.inv()))
        for text in ('Second','4','400'): await self.message(text)
        data=decode_invoice_analysis(await self.inv())
        self.assertEqual(len(data['items']),2)
        self.assertEqual(data['amount_payable'],1692)
        self.assertEqual(data['subtotal_before_discount'],1692)
        await self.click(cb('confirm',await self.inv()))
        await self.message('5')
        inv=await self.inv()
        self.assertEqual(inv['deal_amount'],1692)
        self.assertEqual(inv['reward_amount'],84.6)

    async def test_manual_global_discount_updates_with_new_items_and_can_reset(self):
        from services.invoice_editor import HEADER_FIELDS
        from handlers.admin.invoice_manual import draft_analysis
        await self.start_empty_manual()
        for text in ('Supplier','First','2','100','90','-'): await self.message(text)
        await self.manual_button('headers')
        await self.manual_button(f'field_h_-1_{list(HEADER_FIELDS).index("discount_amount")}')
        await self.message('30')
        data=await self.manual_state().get_data()
        self.assertEqual(draft_analysis(data)['amount_payable'],170)
        await self.manual_button('resume')
        await self.manual_button('next')
        for text in ('Second','1','50','-','-'): await self.message(text)
        data=await self.manual_state().get_data()
        self.assertEqual(draft_analysis(data)['amount_payable'],220)
        self.assertEqual(draft_analysis(data)['calculated_total'],230)
        await self.manual_button('headers')
        await self.manual_button('auto_total')
        data=await self.manual_state().get_data()
        self.assertEqual(draft_analysis(data)['amount_payable'],230)
        self.assertEqual(draft_analysis(data)['discount_amount'],20)

    async def test_manual_discount_survives_reedit_without_losing_last_item(self):
        from services.invoice_editor import HEADER_FIELDS
        from handlers.admin.invoice_manual import draft_analysis
        await self.start_empty_manual()
        for text in ('Supplier','Only','1','100','-','-'): await self.message(text)
        await self.manual_button('headers')
        await self.manual_button(f'field_h_-1_{list(HEADER_FIELDS).index("discount_amount")}')
        await self.message('80')
        await self.manual_button('resume')
        await self.manual_button('edit_item')
        await self.manual_button('headers')
        data=await self.manual_state().get_data()
        self.assertEqual(len(data['items']),1)
        self.assertEqual(draft_analysis(data)['amount_payable'],20)

    async def test_approved_edit_updates_accounting_atomically(self):
        await approve_invoice(self.path,self.ident,180,9,analysis()['items'])
        inv=await self.inv()
        a=apply_change(analysis(),0,'product_name','Новое имя')
        a=apply_change(a,-1,'amount_payable',200)
        await save_invoice_analysis(self.path,self.ident,a,expected_revision=revision(inv))
        inv=await self.inv(); items=await list_invoice_items(self.path,self.ident)
        self.assertEqual(inv['deal_amount'],200); self.assertEqual(inv['reward_amount'],10)
        self.assertEqual(items[0]['product_name'],'Новое имя')

    async def test_recognition_lock_prevents_double_job_and_edits(self):
        from services.invoice_recognition import InvoiceRecognitionError
        inv=await self.inv(); await mark_invoice_analysis_processing(self.path,self.ident)
        with self.assertRaises(InvoiceRecognitionError): await mark_invoice_analysis_processing(self.path,self.ident)
        with self.assertRaises(ValueError): await save_invoice_analysis(self.path,self.ident,analysis(),expected_revision=revision(inv))

    async def test_rich_fallback_without_network(self):
        fake=SimpleNamespace(chat=SimpleNamespace(id=1),bot=AsyncMock())
        fake.bot.side_effect=TelegramNotFound(method=SimpleNamespace(),message='Not Found')
        fake.bot.send_message=AsyncMock()
        from aiogram.types import InlineKeyboardMarkup
        await display(fake,'<b>Card</b>',InlineKeyboardMarkup(inline_keyboard=[]),rich='<p>Card</p>')
        fake.bot.send_message.assert_awaited_once()

if __name__=='__main__': unittest.main()
